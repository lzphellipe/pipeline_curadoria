#!/usr/bin/env python3
"""
Pipeline de limpeza e curadoria quimiotaxonômica de anotações LC-MS/MS.

Etapas
  1. carregar     lê a lista bruta (.xlsx com colunas Número | Componente)
  2. limpar       separa biblioteca, ID do banco e energia de colisão do nome;
                  detecta problemas de qualidade (encoding, espaços, caixa alta...)
  3. padronizar   converte cada nome limpo em um composto canônico (dados/sinonimos.csv)
                  e marca duplicatas
  4. curar        cruza cada composto com a base de conhecimento (dados/base_curadoria.csv):
                  classe química, categoria A–F, grupo isomérico, nível MSI, justificativa
                  e referências. Nome sem correspondência vira categoria P (pendente).
  5. verificar    procura os marcadores esperados da espécie e gera alertas
  6. exportar     planilha .xlsx (LEIA-ME, Resumo, Bruto, Limpo, Curadoria, Referências)
                  e CSVs (limpo, curadoria, pendentes)

Uso
  python pipeline.py dados/Lista_de_componentes_29.09.xlsx -o saida/
  python pipeline.py lista.xlsx -o saida/ --especie "Vanilla planifolia" --familia Orchidaceae

Requisitos: Python 3.9+ e openpyxl  (pip install openpyxl)
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).resolve().parent
DADOS = AQUI / "dados"

# --------------------------------------------------------------------------- configuração
CATEGORIAS = {
    "A": ("Relatado na espécie", "Manter"),
    "B": ("Relatado no gênero ou em outras espécies da família", "Manter"),
    "C": ("Metabólito primário/ubíquo em plantas — plausível", "Manter"),
    "D": ("Natural e biossinteticamente possível, sem registro encontrado na família", "Manter com ressalva — confirmar com padrão"),
    "E": ("Improvável na família (típico de outro táxon, microrganismo ou animal)", "Excluir ou reavaliar"),
    "F": ("Contaminante, sintético, fármaco ou artefato analítico", "Excluir"),
    "P": ("PENDENTE — composto sem curadoria na base de conhecimento", "Curar manualmente"),
}
CORES = {"A": "C6EFCE", "B": "DDEBF7", "C": "E2EFDA", "D": "FFF2CC", "E": "FCE4D6", "F": "F8CBAD", "P": "D9D9D9"}

# Marcadores químicos característicos da espécie. A etapa 5 avisa quando eles não aparecem.
MARCADORES_ESPERADOS = ["Vanillin", "Glucovanillin", "Vanillic acid", "4-Hydroxybenzaldehyde", "4-Hydroxybenzoic acid"]

# Energias de colisão escalonadas, no formato em que algumas bibliotecas exportam
ENERGIAS_ESCALONADAS = {"102040": "10/20/40 (escalonada)", "205060": "20/50/60 (escalonada)"}
BANCOS = {"massbank": "MassBank", "mona": "MoNA", "hmdb": "HMDB"}

# Regras extras de qualidade: (regex aplicada ao nome limpo, descrição do problema)
REGRAS_QUALIDADE = [
    (r"[A-Z]{5,}", "caixa alta inconsistente"),
    (r"\|", "sinônimos concatenados com '|'"),
    (r";\s*\[M", "aduto e fórmula embutidos no nome"),
    (r"\?", "caractere corrompido (encoding)"),
    (r"^[A-Z]{2}\d{5,}$", "código no lugar do nome"),
    (r"^\d+\.\d+[RS]-", "nomenclatura irregular"),
    (r"\bdihy dro", "espaço indevido dentro do nome"),
]


# --------------------------------------------------------------------------- modelo
@dataclass
class Registro:
    numero: int
    original: str
    nome_limpo: str = ""
    biblioteca: str = "Não informada"
    id_banco: str = ""
    energia: str = "Não informada"
    problemas: list[str] = field(default_factory=list)
    composto: str = ""
    duplicata: bool = False


def normalizar_chave(nome: str) -> str:
    return re.sub(r"\s+", " ", nome.strip()).lower()


# --------------------------------------------------------------------------- 1. carregar
def carregar(caminho: Path) -> list[Registro]:
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    regs = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha[0] is None or linha[1] is None:
            continue
        regs.append(Registro(numero=int(linha[0]), original=str(linha[1])))
    if not regs:
        sys.exit("Nenhuma linha encontrada. Esperado: coluna A = Número, coluna B = Componente.")
    return regs


# --------------------------------------------------------------------------- 2. limpar
def limpar(r: Registro) -> None:
    s = r.original
    if s != s.strip():
        r.problemas.append("espaço extra no início/fim")
    s = s.strip()

    m = re.search(r"\s+from (NIST\d+)$", s)
    if m:
        r.biblioteca = m.group(1)
        s = s[: m.start()]
        r.problemas.append("sufixo de biblioteca no nome")

    m = re.match(r"(Massbank|MoNA|HMDB):(\S+)\s+(.*)", s, flags=re.I)
    if m:
        r.biblioteca = BANCOS[m.group(1).lower()]
        r.id_banco = m.group(2)
        s = m.group(3)
        r.problemas.append("ID do banco embutido no nome")

    m = re.search(r"\s*CollisionEnergy:(\d+)$", s)
    if m:
        r.energia = ENERGIAS_ESCALONADAS.get(m.group(1), m.group(1))
        s = s[: m.start()]
        r.problemas.append("energia de colisão embutida no nome")

    m = re.search(r"\s*-\s*([\d.]+)\s*eV$", s)
    if m:
        r.energia = f"{float(m.group(1)):g} eV"
        s = s[: m.start()]
        r.problemas.append("energia de colisão embutida no nome")

    for padrao, descricao in REGRAS_QUALIDADE:
        if re.search(padrao, s):
            r.problemas.append(descricao)

    r.problemas = list(dict.fromkeys(r.problemas))
    r.nome_limpo = s


# --------------------------------------------------------------------------- 3. padronizar
def ler_csv(nome: str) -> list[dict]:
    with open(DADOS / nome, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def padronizar(regs: list[Registro], sinonimos: dict[str, str]) -> None:
    vistos = set()
    for r in regs:
        r.composto = sinonimos.get(normalizar_chave(r.nome_limpo), r.nome_limpo)
        r.duplicata = r.composto in vistos
        vistos.add(r.composto)


# --------------------------------------------------------------------------- 4. curar
def curar(regs: list[Registro], base: dict[str, dict]) -> dict[str, dict]:
    """Retorna um registro por composto único, com a curadoria e as linhas de origem."""
    unicos: dict[str, dict] = {}
    for r in regs:
        if r.composto not in unicos:
            info = base.get(r.composto) or {
                "composto": r.composto, "classe": "A classificar", "categoria": "P",
                "grupo_isomerico": "", "nivel_msi": "2",
                "justificativa": "Não encontrado na base de conhecimento. Pesquisar ocorrência na família "
                                 "(LOTUS, KNApSAcK, literatura) e adicionar a dados/base_curadoria.csv.",
                "referencias": "",
            }
            unicos[r.composto] = {**info, "linhas": []}
        unicos[r.composto]["linhas"].append(r.numero)
    ordem = "ABCDEFP"
    return dict(sorted(unicos.items(), key=lambda kv: (ordem.index(kv[1]["categoria"]), kv[1]["classe"], kv[0].lower())))


# --------------------------------------------------------------------------- 5. verificar
def verificar(unicos: dict[str, dict], regs: list[Registro], especie: str) -> list[str]:
    alertas = []
    nomes = {k.lower() for k in unicos}
    ausentes = [m for m in MARCADORES_ESPERADOS if m.lower() not in nomes]
    if ausentes:
        alertas.append(
            f"Ausências notáveis: {', '.join(ausentes)}, marcadores característicos de {especie}, não aparecem na lista. "
            "Verificar o score de corte, outra exportação ou o tecido/extrato analisado.")
    if any("tetrapirrol" in d["classe"].lower() for d in unicos.values()):
        alertas.append("Há derivados de clorofila (feofórbido/feofitina), o que indica tecido verde (folha, caule ou vagem verde).")
    pend = [k for k, d in unicos.items() if d["categoria"] == "P"]
    if pend:
        alertas.append(f"{len(pend)} composto(s) sem curadoria (categoria P). Ver pendentes.csv.")
    alertas.append("Todas as anotações vêm de similaridade espectral em biblioteca, portanto são no máximo nível 2 do MSI "
                   "(Sumner et al., 2007). Nível 1 exige padrão analítico com o mesmo tempo de retenção e o mesmo MS/MS.")
    alertas.append("Recomenda-se rodar um branco de extração (solvente + vial) para confirmar os contaminantes da categoria F.")
    alertas.append("'Sem registro encontrado' não equivale a 'ausente na família'. Antes de excluir compostos D ou E, "
                   "consultar LOTUS (lotus.naturalproducts.net) e KNApSAcK.")
    return alertas


# --------------------------------------------------------------------------- 6. exportar
FONTE = "Arial"
F_CORPO = Font(name=FONTE, size=10)
F_NEGRITO = Font(name=FONTE, size=10, bold=True)
F_CAB = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
F_TIT = Font(name=FONTE, size=14, bold=True, color="1F3A5F")
F_SEC = Font(name=FONTE, size=11, bold=True, color="1F3A5F")
P_CAB = PatternFill("solid", fgColor="1F3A5F")
QUEBRA = Alignment(wrap_text=True, vertical="top")
TOPO = Alignment(vertical="top")


def _cabecalho(ws, linha, colunas, larguras):
    for i, (c, w) in enumerate(zip(colunas, larguras), 1):
        cel = ws.cell(linha, i, c)
        cel.fill, cel.font = P_CAB, F_CAB
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[linha].height = 32


def _cores_categoria(ws, intervalo, coluna):
    for k, cor in CORES.items():
        ws.conditional_formatting.add(intervalo, FormulaRule(formula=[f'${coluna}2="{k}"'], fill=PatternFill("solid", fgColor=cor)))


def exportar_xlsx(regs, unicos, refs, alertas, saida: Path, especie, familia, nome_entrada):
    wb = openpyxl.Workbook()
    n = len(regs)
    ult_l = n + 1                      # última linha da aba 2_Limpo
    ult_c = len(unicos) + 1            # última linha da aba 3_Curadoria
    cats = list(CATEGORIAS)
    r0, r1 = 5, 4 + len(cats)          # linhas da tabela de categorias no Resumo
    tot = r1 + 1

    # ---------------- LEIA-ME
    ws = wb.active
    ws.title = "LEIA-ME"
    for col, w in zip("ABCD", [4, 34, 70, 30]):
        ws.column_dimensions[col].width = w
    ws["B2"] = f"Curadoria de anotações LC-MS/MS — {especie} ({familia})"
    ws["B2"].font = F_TIT
    ws["B3"] = f"Fonte: {nome_entrada} ({n} linhas). Gerado por pipeline.py."
    ws["B3"].font = F_CORPO
    lin = 5
    ws.cell(lin, 2, "Etapas do pipeline").font = F_SEC
    lin += 1
    etapas = [
        ("1. Carregar", "Aba 1_Bruto mantém os dados originais intactos."),
        ("2. Limpar", "Aba 2_Limpo: separa do nome a biblioteca, o ID do banco e a energia de colisão, e registra os problemas de qualidade de cada linha."),
        ("3. Padronizar", "Cada nome limpo é convertido em um composto canônico (dados/sinonimos.csv). Linhas repetidas do mesmo composto são marcadas como duplicatas."),
        ("4. Curar", "Aba 3_Curadoria: cada composto único recebe classe, categoria A–F, grupo isomérico, nível MSI, justificativa e referências (dados/base_curadoria.csv)."),
        ("5. Verificar", "Procura os marcadores esperados da espécie e gera os alertas listados abaixo."),
        ("6. Resumir", "Aba 4_Resumo: contagens por categoria e classe, calculadas por fórmula. Se você mudar uma categoria na 3_Curadoria, o resumo se atualiza."),
    ]
    for a, b in etapas:
        ws.cell(lin, 2, a).font = F_NEGRITO
        c = ws.cell(lin, 3, b)
        c.font, c.alignment = F_CORPO, QUEBRA
        ws.row_dimensions[lin].height = 15 * (len(b) // 85 + 1)
        lin += 1
    lin += 1
    ws.cell(lin, 2, "Categorias de plausibilidade").font = F_SEC
    lin += 1
    for i, h in enumerate(["Categoria", "Significado", "Ação recomendada"]):
        c = ws.cell(lin, 2 + i, h)
        c.fill, c.font = P_CAB, F_CAB
    lin += 1
    for k, (desc, acao) in CATEGORIAS.items():
        for j, v in enumerate([k, desc, acao]):
            c = ws.cell(lin, 2 + j, v)
            c.font = F_NEGRITO if j == 0 else F_CORPO
            c.fill = PatternFill("solid", fgColor=CORES[k])
        lin += 1
    lin += 1
    ws.cell(lin, 2, "Alertas e observações").font = F_SEC
    lin += 1
    for t in alertas:
        c = ws.cell(lin, 2, "• " + t)
        c.font, c.alignment = F_CORPO, QUEBRA
        ws.merge_cells(start_row=lin, start_column=2, end_row=lin, end_column=4)
        ws.row_dimensions[lin].height = 15 * (len(t) // 120 + 1)
        lin += 1
    ws.sheet_view.showGridLines = False

    # ---------------- 4_Resumo (criada antes para ficar na 2ª posição)
    wr = wb.create_sheet("4_Resumo")

    # ---------------- 1_Bruto
    wbr = wb.create_sheet("1_Bruto")
    _cabecalho(wbr, 1, ["Número", "Componente (original)"], [10, 110])
    for i, r in enumerate(regs, 2):
        wbr.cell(i, 1, r.numero).font = F_CORPO
        wbr.cell(i, 2, r.original).font = F_CORPO
    wbr.freeze_panes = "A2"

    # ---------------- 2_Limpo
    wl = wb.create_sheet("2_Limpo")
    _cabecalho(wl, 1, ["Nº original", "Nome original", "Composto padronizado", "Biblioteca", "ID no banco",
                       "Energia de colisão", "Problemas de qualidade encontrados", "Nº de linhas do mesmo composto",
                       "Duplicata?", "Categoria (A–F)", "Ação recomendada"],
               [9, 55, 42, 12, 12, 18, 42, 12, 11, 11, 30])
    for i, r in enumerate(regs, 2):
        valores = [r.numero, r.original.strip(), r.composto, r.biblioteca, r.id_banco, r.energia,
                   "; ".join(r.problemas) or "—"]
        for j, v in enumerate(valores, 1):
            c = wl.cell(i, j, v)
            c.font, c.alignment = F_CORPO, (QUEBRA if j in (2, 3, 7) else TOPO)
        wl.cell(i, 8, f"=COUNTIF($C$2:$C${ult_l},C{i})")
        wl.cell(i, 9, f'=IF(COUNTIF($C$2:C{i},C{i})>1,"Sim","Não")')
        wl.cell(i, 10, f"=INDEX('3_Curadoria'!$F$2:$F${ult_c},MATCH(C{i},'3_Curadoria'!$B$2:$B${ult_c},0))")
        wl.cell(i, 11, f"=INDEX('3_Curadoria'!$H$2:$H${ult_c},MATCH(C{i},'3_Curadoria'!$B$2:$B${ult_c},0))")
        for j in range(8, 12):
            wl.cell(i, j).font = F_CORPO
    wl.freeze_panes = "C2"
    wl.auto_filter.ref = f"A1:K{ult_l}"
    _cores_categoria(wl, f"J2:J{ult_l}", "J")
    wl.conditional_formatting.add(f"I2:I{ult_l}", FormulaRule(formula=['$I2="Sim"'], font=Font(name=FONTE, color="9C0006", bold=True)))

    # ---------------- 3_Curadoria
    wc = wb.create_sheet("3_Curadoria")
    _cabecalho(wc, 1, ["ID", "Composto padronizado", "Classe química", "Nº de linhas na lista", "Linhas originais",
                       "Categoria (A–F)", "Plausibilidade", "Ação recomendada", "Grupo isomérico / ressalva de anotação",
                       "Nível MSI sugerido", "Justificativa", "Referências"],
               [6, 40, 26, 10, 16, 10, 34, 24, 30, 9, 70, 12])
    for i, (k, d) in enumerate(unicos.items(), 2):
        valores = {1: i - 1, 2: k, 3: d["classe"], 5: ", ".join(map(str, d["linhas"])), 6: d["categoria"],
                   9: d["grupo_isomerico"] or "—", 10: d["nivel_msi"], 11: d["justificativa"],
                   12: d["referencias"].replace(";", "; ") or "—"}
        for j, v in valores.items():
            c = wc.cell(i, j, v)
            c.font, c.alignment = F_CORPO, (QUEBRA if j in (2, 3, 5, 9, 11) else TOPO)
        wc.cell(i, 4, f"=COUNTIF('2_Limpo'!$C$2:$C${ult_l},B{i})")
        wc.cell(i, 7, f"=INDEX('4_Resumo'!$C${r0}:$C${r1},MATCH(F{i},'4_Resumo'!$B${r0}:$B${r1},0))")
        wc.cell(i, 8, f"=INDEX('4_Resumo'!$D${r0}:$D${r1},MATCH(F{i},'4_Resumo'!$B${r0}:$B${r1},0))")
        for j in (4, 7, 8):
            wc.cell(i, j).font, wc.cell(i, j).alignment = F_CORPO, QUEBRA
    wc.freeze_panes = "C2"
    wc.auto_filter.ref = f"A1:L{ult_c}"
    _cores_categoria(wc, f"F2:F{ult_c}", "F")
    dv = DataValidation(type="list", formula1='"' + ",".join(cats) + '"')
    wc.add_data_validation(dv)
    dv.add(f"F2:F{ult_c}")

    # ---------------- 4_Resumo (conteúdo)
    wr.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFG", [12, 58, 36, 16, 18, 14]):
        wr.column_dimensions[col].width = w
    wr["B2"] = "Resumo da curadoria"
    wr["B2"].font = F_TIT
    for j, h in enumerate(["Categoria", "Plausibilidade", "Ação recomendada", "Compostos únicos",
                           "Linhas da lista original", "% dos únicos"], 2):
        c = wr.cell(4, j, h)
        c.fill, c.font, c.alignment = P_CAB, F_CAB, Alignment(wrap_text=True, vertical="center")
    wr.row_dimensions[4].height = 30
    for i, (k, (desc, acao)) in enumerate(CATEGORIAS.items(), r0):
        wr.cell(i, 2, k)
        wr.cell(i, 3, desc)
        wr.cell(i, 4, acao)
        wr.cell(i, 5, f"=COUNTIF('3_Curadoria'!$F$2:$F${ult_c},B{i})")
        wr.cell(i, 6, f"=SUMIF('3_Curadoria'!$F$2:$F${ult_c},B{i},'3_Curadoria'!$D$2:$D${ult_c})")
        wr.cell(i, 7, f"=IF($E${tot}=0,0,E{i}/$E${tot})").number_format = "0.0%"
        for col in range(2, 8):
            c = wr.cell(i, col)
            c.font, c.alignment, c.fill = F_CORPO, QUEBRA, PatternFill("solid", fgColor=CORES[k])
    wr.cell(tot, 2, "Total").font = F_NEGRITO
    for col, letra in ((5, "E"), (6, "F"), (7, "G")):
        c = wr.cell(tot, col, f"=SUM({letra}{r0}:{letra}{r1})")
        c.font = F_NEGRITO
    wr.cell(tot, 7).number_format = "0.0%"

    lin = tot + 2
    wr.cell(lin, 2, "Indicadores do ciclo de limpeza").font = F_SEC
    idx = {k: r0 + i for i, k in enumerate(cats)}
    indicadores = [
        ("Linhas na lista original", f"=COUNTA('1_Bruto'!A2:A{ult_l})"),
        ("Compostos únicos após padronização", f"=COUNTA('3_Curadoria'!B2:B{ult_c})"),
        ("Linhas redundantes (duplicatas)", f"=COUNTIF('2_Limpo'!I2:I{ult_l},\"Sim\")"),
        ("Linhas com algum problema de qualidade", f"=COUNTIF('2_Limpo'!G2:G{ult_l},\"<>—\")"),
        ("Compostos mantidos (A + B + C)", f"=E{idx['A']}+E{idx['B']}+E{idx['C']}"),
        ("Compostos a confirmar (D)", f"=E{idx['D']}"),
        ("Compostos a excluir/reavaliar (E + F)", f"=E{idx['E']}+E{idx['F']}"),
        ("Compostos pendentes de curadoria (P)", f"=E{idx['P']}"),
        ("Compostos em grupos isoméricos", f"=COUNTIF('3_Curadoria'!I2:I{ult_c},\"<>—\")"),
    ]
    for i, (a, f) in enumerate(indicadores, lin + 1):
        wr.cell(i, 3, a).font = F_CORPO
        wr.cell(i, 5, f).font = F_NEGRITO

    lin = lin + len(indicadores) + 2
    wr.cell(lin, 2, "Compostos únicos por classe química e categoria").font = F_SEC
    for j, h in enumerate(["Classe química"] + cats + ["Total"], 3):
        c = wr.cell(lin + 1, j, h)
        c.fill, c.font = P_CAB, F_CAB
    ult_cat_col = get_column_letter(3 + len(cats))
    for i, cl in enumerate(sorted({d["classe"] for d in unicos.values()}), lin + 2):
        wr.cell(i, 3, cl).font = F_CORPO
        for j, k in enumerate(cats, 4):
            wr.cell(i, j, f"=COUNTIFS('3_Curadoria'!$C$2:$C${ult_c},$C{i},'3_Curadoria'!$F$2:$F${ult_c},\"{k}\")").font = F_CORPO
        wr.cell(i, 4 + len(cats), f"=SUM(D{i}:{ult_cat_col}{i})").font = F_NEGRITO
    for j in range(4, 5 + len(cats)):
        wr.column_dimensions[get_column_letter(j)].width = max(wr.column_dimensions[get_column_letter(j)].width or 0, 8)

    graf = BarChart()
    graf.type, graf.style, graf.title, graf.legend = "bar", 10, "Compostos únicos por categoria", None
    graf.add_data(Reference(wr, min_col=5, min_row=4, max_row=r1), titles_from_data=True)
    graf.set_categories(Reference(wr, min_col=2, min_row=r0, max_row=r1))
    graf.height, graf.width = 7, 12
    wr.add_chart(graf, "I4")
    wr.sheet_view.showGridLines = False

    # ---------------- 5_Referências
    wf = wb.create_sheet("5_Referências")
    _cabecalho(wf, 1, ["Código", "Referência", "Link"], [9, 110, 60])
    for i, ref in enumerate(refs, 2):
        wf.cell(i, 1, ref["codigo"]).font = F_CORPO
        c = wf.cell(i, 2, ref["referencia"])
        c.font, c.alignment = F_CORPO, QUEBRA
        c = wf.cell(i, 3, ref["link"])
        c.hyperlink = ref["link"]
        c.font = Font(name=FONTE, size=10, color="0563C1", underline="single")

    wb.calculation.fullCalcOnLoad = True   # Excel/LibreOffice calculam as fórmulas ao abrir
    wb.save(saida)


def exportar_csvs(regs, unicos, pasta: Path):
    with open(pasta / "2_limpo.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["numero", "nome_original", "composto", "biblioteca", "id_banco", "energia_colisao", "problemas", "duplicata", "categoria"])
        for r in regs:
            w.writerow([r.numero, r.original.strip(), r.composto, r.biblioteca, r.id_banco, r.energia,
                        "; ".join(r.problemas), "Sim" if r.duplicata else "Não", unicos[r.composto]["categoria"]])
    with open(pasta / "3_curadoria.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["composto", "classe", "n_linhas", "linhas", "categoria", "grupo_isomerico", "nivel_msi", "justificativa", "referencias"])
        for k, d in unicos.items():
            w.writerow([k, d["classe"], len(d["linhas"]), ",".join(map(str, d["linhas"])), d["categoria"],
                        d["grupo_isomerico"], d["nivel_msi"], d["justificativa"], d["referencias"]])
    pend = [(k, d) for k, d in unicos.items() if d["categoria"] == "P"]
    with open(pasta / "pendentes.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["composto", "linhas"])
        for k, d in pend:
            w.writerow([k, ",".join(map(str, d["linhas"]))])


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Limpeza e curadoria quimiotaxonômica de anotações LC-MS/MS")
    ap.add_argument("entrada", type=Path, help="planilha .xlsx (coluna A = Número, coluna B = Componente)")
    ap.add_argument("-o", "--saida", type=Path, default=Path("saida"), help="pasta de saída (padrão: saida/)")
    ap.add_argument("--especie", default="Vanilla planifolia")
    ap.add_argument("--familia", default="Orchidaceae")
    args = ap.parse_args()
    args.saida.mkdir(parents=True, exist_ok=True)

    print(f"[1/6] Carregando {args.entrada}")
    regs = carregar(args.entrada)

    print("[2/6] Limpando nomes")
    for r in regs:
        limpar(r)

    print("[3/6] Padronizando e deduplicando")
    sinonimos = {normalizar_chave(s["nome_limpo"]): s["composto"] for s in ler_csv("sinonimos.csv")}
    padronizar(regs, sinonimos)

    print("[4/6] Curando com a base de conhecimento")
    base = {b["composto"]: b for b in ler_csv("base_curadoria.csv")}
    unicos = curar(regs, base)

    print("[5/6] Verificando marcadores e gerando alertas")
    alertas = verificar(unicos, regs, args.especie)

    print("[6/6] Exportando")
    nome = f"Curadoria_compostos_{args.especie.replace(' ', '_')}.xlsx"
    exportar_xlsx(regs, unicos, ler_csv("referencias.csv"), alertas, args.saida / nome,
                  args.especie, args.familia, args.entrada.name)
    exportar_csvs(regs, unicos, args.saida)

    cont = Counter(d["categoria"] for d in unicos.values())
    print("\n================ RESUMO ================")
    print(f"Linhas na entrada ..................... {len(regs)}")
    print(f"Compostos únicos ...................... {len(unicos)}")
    print(f"Linhas duplicadas ..................... {sum(r.duplicata for r in regs)}")
    print(f"Linhas com problema de qualidade ...... {sum(bool(r.problemas) for r in regs)}")
    for k in CATEGORIAS:
        if cont[k]:
            print(f"  {k} – {CATEGORIAS[k][0][:58]:<58} {cont[k]:>3}")
    print("\nAlertas:")
    for a in alertas:
        print("  •", a)
    print(f"\nArquivos em {args.saida.resolve()}/")


if __name__ == "__main__":
    main()
